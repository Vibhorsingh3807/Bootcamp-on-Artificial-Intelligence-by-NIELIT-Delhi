import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_demo_docx(filepath):
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    navy = RGBColor(15, 23, 42)
    accent = RGBColor(2, 132, 199)
    dark_text = RGBColor(30, 41, 59)

    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("EchoSign AI — Live Stage Demo & Failsafe Sign Guide")
    r.font.name = "Arial"
    r.font.size = Pt(19)
    r.font.bold = True
    r.font.color.rgb = navy

    # Subtitle
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("Foolproof Sentences, Hand Gestures, NLP Reconstructions & Stage Presentation Playbook")
    r2.font.name = "Arial"
    r2.font.size = Pt(10.5)
    r2.font.bold = True
    r2.font.color.rgb = accent

    # Meta
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_after = Pt(14)
    r3 = p3.add_run("Bootcamp on Artificial Intelligence by NIELIT Delhi  |  Presentation Handout")
    r3.font.name = "Arial"
    r3.font.size = Pt(9.5)
    r3.font.italic = True
    r3.font.color.rgb = RGBColor(100, 116, 139)

    # 1. Golden Rules
    h1 = doc.add_heading(level=1)
    h1.add_run("1. The 3 Stage Golden Rules (Zero Failure Guarantee)").font.color.rgb = navy
    doc.add_paragraph("1. Camera Distance: Sit 2.5 to 3 feet from the webcam so Face and Chest are framed clearly.\n"
                      "2. Room Lighting: Front-lit hands and face; avoid strong backlights.\n"
                      "3. Hold for 1.5 Seconds: Hold each sign steady until the HUD locks and speaks aloud.")

    # 2. Healthcare & Medical Urgency
    h2 = doc.add_heading(level=1)
    h2.add_run("2. Healthcare & Medical Urgency (Hospital / Pain / Doctor)").font.color.rgb = navy

    t_med = doc.add_table(rows=1, cols=4)
    t_med.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, title in enumerate(["Gesture / Combo", "Physical Hand Movement", "AI Spoken Output Sentence", "Emotion & Urgency"]):
        c = t_med.rows[0].cells[i]
        c.text = title
        set_cell_background(c, "0F172A")
        set_cell_margins(c, 100, 100, 120, 120)
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    med_rows = [
        ("PAIN + DOCTOR\n(Emergency)", "1. Point index finger to chest/stomach (PAIN).\n2. Tap 2 fingers on opposite wrist pulse (DOCTOR).", "'I am experiencing pain and would like to see a medical professional.'", "Medical Attention\nUrgency: 0.50 (Moderate)"),
        ("PLEASE + DOCTOR\n(Care Request)", "1. Flat open hand across chest (PLEASE).\n2. Tap opposite wrist or neck pulse (DOCTOR).", "'Could you please assist me in seeing a doctor?'", "Care Requested\nUrgency: 0.40"),
        ("DOCTOR\n(Single Sign)", "Index & Middle fingers tap opposite wrist pulse or neck pulse.", "'Could you please connect me with a doctor or medical professional?'", "Medical Attention\nUrgency: 0.40"),
        ("PAIN\n(Single Sign)", "Index finger points directly towards chest/stomach area.", "'I am experiencing some physical discomfort or pain.'", "Care Requested\nUrgency: 0.40")
    ]
    for idx, (g, m, s, e) in enumerate(med_rows):
        row = t_med.add_row()
        for i, val in enumerate([g, m, s, e]):
            c = row.cells[i]
            c.text = val
            set_cell_background(c, "F8FAFC" if idx % 2 == 1 else "FFFFFF")
            set_cell_margins(c, 90, 90, 110, 110)

    # 3. Needs (Food & Water)
    h3 = doc.add_heading(level=1)
    h3.add_run("3. Basic Needs (Food, Water, Hunger)").font.color.rgb = navy
    t_need = doc.add_table(rows=1, cols=4)
    t_need.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, title in enumerate(["Gesture / Combo", "Physical Hand Movement", "AI Spoken Output Sentence", "Emotion & Urgency"]):
        c = t_need.rows[0].cells[i]
        c.text = title
        set_cell_background(c, "0F172A")
        set_cell_margins(c, 100, 100, 120, 120)
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    need_rows = [
        ("PLEASE + WATER\n(Drink Request)", "1. Flat hand on chest (PLEASE).\n2. 'W' shape touches lips/chin (WATER).", "'Could I please get some drinking water?'", "Polite Request\nUrgency: 0.20"),
        ("PLEASE + FOOD\n(Meal Request)", "1. Flat hand on chest (PLEASE).\n2. Pinch 5 fingertips together touching lips (FOOD).", "'Could I please have some food?'", "Polite Request\nUrgency: 0.20"),
        ("WATER + FOOD\n(Complete Need)", "1. 'W' shape at lips (WATER).\n2. Pinch cluster touching lips (FOOD).", "'Could I please get some drinking water and food?'", "Polite Request\nUrgency: 0.20"),
        ("FOOD\n(Single Sign)", "Pinch all 5 fingertips together and touch lips/mouth repeatedly.", "'I am feeling hungry, could you please provide some food?'", "Request / Need\nUrgency: 0.20"),
        ("WATER\n(Single Sign)", "'W' shape (Index, Middle, Ring up) tapped near chin/lips.", "'Could I please have a glass of drinking water?'", "Request / Need\nUrgency: 0.20")
    ]
    for idx, (g, m, s, e) in enumerate(need_rows):
        row = t_need.add_row()
        for i, val in enumerate([g, m, s, e]):
            c = row.cells[i]
            c.text = val
            set_cell_background(c, "F8FAFC" if idx % 2 == 1 else "FFFFFF")
            set_cell_margins(c, 90, 90, 110, 110)

    doc.add_page_break()

    # 4. Greetings, Gratitude & Social
    h4 = doc.add_heading(level=1)
    h4.add_run("4. Greetings, Gratitude & Confirmation (Namaste / OK / Peace)").font.color.rgb = navy
    t_soc = doc.add_table(rows=1, cols=4)
    t_soc.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, title in enumerate(["Gesture / Combo", "Physical Hand Movement", "AI Spoken Output Sentence", "Emotion & Urgency"]):
        c = t_soc.rows[0].cells[i]
        c.text = title
        set_cell_background(c, "0F172A")
        set_cell_margins(c, 100, 100, 120, 120)
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    soc_rows = [
        ("NAMASTE\n(Dual Hand)", "Bring both palms together at chest level with fingers pointing upward.", "'Namaste! Welcome and heartfelt greetings to you.'", "Warm & Grateful\nConfidence: 0.98"),
        ("HELLO\n(Single Hand)", "Open 5 fingers upright facing camera, wave near temple height.", "'Hello! Wishing you a peaceful and pleasant day.'", "Conversational\nConfidence: 0.97"),
        ("THANK YOU\n(Single Sign)", "Fingers of flat hand touch lips/chin and extend outward towards camera.", "'Thank you so much! Truly appreciate your kind support.'", "Warm & Grateful"),
        ("OK / GOOD\n(👌 Sign)", "Thumb and Index finger tips touch forming a circle; Middle, Ring, Pinky upright.", "'All is well, everything is completely okay.'", "Positive\nConfidence: 0.96"),
        ("YES\n(Thumbs Up 👍)", "Closed fist with thumb upright pointing to the ceiling.", "'Yes, absolutely correct.'", "Agreement\nConfidence: 0.95"),
        ("YES + THANK YOU", "1. Thumbs Up (YES).\n2. Flat hand extending from lips (THANK YOU).", "'Yes, that is correct! Thank you so much.'", "Warm & Grateful"),
        ("NO + THANK YOU", "1. Index & Middle finger snap down (NO).\n2. Flat hand lips outward (THANK YOU).", "'No, thank you, but I appreciate your offer.'", "Polite Refusal"),
        ("I LOVE YOU\n(🤟 Sign)", "Thumb, Index, and Pinky fingers open simultaneously (Middle & Ring curled).", "'I love and appreciate you deeply!'", "Social / Affection\nConfidence: 0.96"),
        ("PEACE\n(✌️ V-Sign)", "Index and Middle finger spread wide in 'V' shape (Ring, Pinky, Thumb tucked).", "'Peace, calmness, and goodwill to all!'", "Calmness\nConfidence: 0.96")
    ]
    for idx, (g, m, s, e) in enumerate(soc_rows):
        row = t_soc.add_row()
        for i, val in enumerate([g, m, s, e]):
            c = row.cells[i]
            c.text = val
            set_cell_background(c, "F8FAFC" if idx % 2 == 1 else "FFFFFF")
            set_cell_margins(c, 90, 90, 110, 110)

    # 5. Reverse Channel & Stage Script
    h5 = doc.add_heading(level=1)
    h5.add_run("5. Reverse Channel (Voice → Sign) Demonstration Lines").font.color.rgb = navy
    doc.add_paragraph("• Hospital Urgency: 'Please call a doctor, I have pain.' → Cards: PLEASE, DOCTOR, PAIN\n"
                      "• Meal Hospitality: 'Are you hungry? Do you need food and water?' → Cards: FOOD, WATER\n"
                      "• Social Respect: 'Namaste, thank you for coming today.' → Cards: NAMASTE, THANK YOU\n"
                      "• Casual Check: 'Hello, is everything okay with you?' → Cards: HELLO, OK")

    h6 = doc.add_heading(level=1)
    h6.add_run("6. Minute-by-Minute Stage Presentation Script").font.color.rgb = navy
    doc.add_paragraph("00:00 - 00:45 | Introduction: Introduce 466M+ hearing impaired problem and EchoSign two-way bridge.\n"
                      "00:45 - 01:45 | Channel A (Deaf → Hearing): Perform dual-hand Namaste, then PLEASE + WATER for NLP sentence reconstruction.\n"
                      "01:45 - 02:45 | Emergency & Social: Perform PAIN + DOCTOR and OK/I Love You signs.\n"
                      "02:45 - 03:30 | Channel B & GPU: Speak 'Please call a doctor, I have pain' in Voice-to-Sign tab. Then switch to Alphabet mode for L and V on RTX 4060 GPU.\n"
                      "03:30 - 04:00 | Conclusion: Invite evaluators' questions.")

    doc.save(filepath)
    print(f"Successfully generated DOCX: {filepath}")

if __name__ == "__main__":
    out_docx = "EchoSign_Live_Demo_and_Sign_Language_Guide.docx"
    create_demo_docx(out_docx)
